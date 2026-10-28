"""Build forms.js (the questions of all three forms, in every language) from the team's Excel files.

Put the files in the translations/ folder (it is not uploaded to GitHub), then run:
    python tools/import_excel.py

- The English file (name contains "English") is the master: question numbers, types, conditions, options.
- Each other file is one language, found from its file name (Marathi, Hindi, Tamil, Telugu, Odia, Kannada, Punjabi).
  It must have the same sheets and rows as the English file (same layout as "1MV Survey Tools _ English.xlsx").
- Anything missing in a language falls back to English in the form.
- The Sheet columns are the English question text, so translations never change the Sheet.
"""
import json, pathlib, re, sys
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / "translations"
OUT = ROOT / "forms.js"

LANG_BY_NAME = {"english": "en", "hindi": "hi", "marathi": "mr", "tamil": "ta", "telugu": "te",
                "odia": "or", "odiya": "or", "oriya": "or", "kannada": "kn", "punjabi": "pa"}

# Sheet name -> form key, and where things are in that sheet (column letters, 1-based index).
FORMS = {
    "youth":  {"sheet": "Youth - Survey",          "type": 4, "required": None, "parent": 5, "when": 6, "opt1": 7, "optN": 16},
    "parent": {"sheet": "Parent - Survey",         "type": 4, "required": None, "parent": 5, "when": 6, "opt1": 7, "optN": 16},
    "public": {"sheet": "General Public - Survey", "type": 5, "required": 4,    "parent": None, "when": None, "opt1": 6, "optN": 16},
}
TYPES = {"single select": "choice", "multiselect": "multi", "multi select": "multi", "dropdown": "dropdown", "drop down": "dropdown",
         "text": "short", "numerical": "number", "number": "number"}


def lang_of(path):
    name = path.stem.lower()
    for word, code in LANG_BY_NAME.items():
        if word in name:
            return code
    return None


def s(v):
    if v is None:
        return ""
    if isinstance(v, float):
        v = repr(v)
    return str(v).replace("\r", "").strip()


def qno(v):
    """0.1 -> '0.1', 2.0 / '2.0' -> '2', '8A' -> '8A'."""
    t = s(v)
    if re.fullmatch(r"\d+\.0", t):
        t = t[:-2]
    return t


def rows_of(ws):
    """Data rows: from row 3 on, every row with text in column C."""
    out = []
    for r in range(3, ws.max_row + 1):
        text = s(ws.cell(r, 3).value)
        if text:
            out.append(r)
    return out


def split_lead(text):
    """'Thank you ... \\n\\nWhat do you want to be?' -> ('Thank you ...', 'What do you want to be?')."""
    parts = [p.strip() for p in re.split(r"\n\s*\n|\n", text) if p.strip()]
    if len(parts) < 2:
        return "", text.strip()
    return "\n".join(parts[:-1]), parts[-1]


def read_form(ws, lay):
    rows = rows_of(ws)
    items, opening, closing, section = [], "", "", ""
    for r in rows:
        no = qno(ws.cell(r, 1).value)
        sec = s(ws.cell(r, 2).value)
        text = s(ws.cell(r, 3).value)
        if sec:
            section = sec
        if not no:
            low = sec.lower()
            if low in ("opening", "प्रस्तावना") or (not items and not opening):
                opening = text
            else:
                closing = text
            continue
        lead, label = split_lead(text)
        opts = [s(ws.cell(r, c).value) for c in range(lay["opt1"], lay["optN"] + 1)]
        opts = [o for o in opts if o]
        items.append({
            "no": no, "section": section, "lead": lead, "label": label,
            "type": s(ws.cell(r, lay["type"]).value).lower(),
            "required": s(ws.cell(r, lay["required"]).value).lower() if lay["required"] else "",
            "parent": qno(ws.cell(r, lay["parent"]).value) if lay["parent"] else "",
            "when": s(ws.cell(r, lay["when"]).value) if lay["when"] else "",
            "options": opts,
        })
    return {"opening": opening, "closing": closing, "items": items}


def item_id(no):
    return "Q" + re.sub(r"[^0-9A-Za-z]", "_", no)


def main():
    files = sorted(p for p in SRC.glob("*.xlsx") if not p.name.startswith("~$"))
    master = [p for p in files if lang_of(p) == "en"]
    if len(master) != 1:
        sys.exit(f"Need exactly one English file in {SRC} (found {len(master)}).")
    langs = {}
    for p in files:
        code = lang_of(p)
        if code and code != "en":
            if code in langs:
                sys.exit(f"Two files for language {code}: {langs[code].name} and {p.name}")
            langs[code] = p
        elif not code:
            print(f"! skipped {p.name}: language not found in the file name")

    en_wb = openpyxl.load_workbook(master[0], data_only=True)
    tr_wbs = {code: openpyxl.load_workbook(p, data_only=True) for code, p in langs.items()}
    warnings, forms = [], {}

    for key, lay in FORMS.items():
        en = read_form(en_wb[lay["sheet"]], lay)
        tr = {}
        for code, wb in tr_wbs.items():
            if lay["sheet"] not in wb.sheetnames:
                warnings.append(f"{code}: no sheet '{lay['sheet']}' (English is used)")
                continue
            tr[code] = read_form(wb[lay["sheet"]], lay)

        def text_of(field, i=None, opt=None):
            out = {}
            for code, f in [("en", en)] + list(tr.items()):
                if i is None:
                    v = f[field]
                else:
                    if i >= len(f["items"]) or f["items"][i]["no"] != en["items"][i]["no"]:
                        continue
                    it = f["items"][i]
                    if opt is not None:
                        if len(it["options"]) != len(en["items"][i]["options"]):
                            continue
                        v = it["options"][opt]
                    else:
                        v = it[field]
                if v:
                    out[code] = v
            return out

        for code, f in tr.items():
            if len(f["items"]) != len(en["items"]):
                warnings.append(f"{code} {key}: {len(f['items'])} questions, English has {len(en['items'])}")
            for i, it in enumerate(en["items"]):
                if i < len(f["items"]) and f["items"][i]["no"] != it["no"]:
                    warnings.append(f"{code} {key}: row {i + 1} is Q{f['items'][i]['no']}, English is Q{it['no']} (English used)")
                elif i < len(f["items"]) and it["type"] != "dropdown" and len(f["items"][i]["options"]) != len(it["options"]):
                    warnings.append(f"{code} {key} Q{it['no']}: {len(f['items'][i]['options'])} options, English has {len(it['options'])} (English options used)")

        ids = {it["no"]: item_id(it["no"]) for it in en["items"]}
        sections, items = [], []
        for i, it in enumerate(en["items"]):
            typ = TYPES.get(it["type"].strip())
            if not typ:
                sys.exit(f"{key} Q{it['no']}: unknown response type '{it['type']}'")
            low = it["label"].lower()
            source = ""
            if typ == "dropdown" and (not it["options"] or it["options"][0].lower().startswith("drop")):
                source = "state" if "state" in low else "district" if "district" in low else ""
                if not source:
                    sys.exit(f"{key} Q{it['no']}: dropdown without options")
            sec_id = re.sub(r"[^a-z0-9]+", "-", it["section"].lower()).strip("-")
            if sec_id not in [x["id"] for x in sections]:
                sections.append({"id": sec_id, "title": {c: f["items"][i]["section"] for c, f in [("en", en)] + list(tr.items())
                                                         if i < len(f["items"]) and f["items"][i]["no"] == it["no"] and f["items"][i]["section"]}})
            q = {"id": ids[it["no"]], "no": it["no"], "section": sec_id, "type": typ,
                 "label": text_of("label", i)}
            lead = text_of("lead", i)
            if lead:
                q["lead"] = lead
            if source:
                q["source"] = source
            elif typ in ("choice", "multi", "dropdown"):
                q["options"] = [dict(v=f"O{j + 1}", **text_of(None, i, j)) for j in range(len(it["options"]))]
            if it["required"]:
                q["required"] = it["required"].startswith("y")
            else:
                q["required"] = typ != "short"          # no column in the sheet: choices required, text optional
            if typ == "short":
                q["probe"] = True
            if typ == "number":
                q["min"], q["max"] = 0, 100000000
            if it["parent"]:
                if it["parent"] not in ids:
                    sys.exit(f"{key} Q{it['no']}: condition on unknown question {it['parent']}")
                codes = [c.strip().upper() for c in re.split(r"[,\s]+", it["when"]) if c.strip()]
                if not codes or any(not re.fullmatch(r"O\d+", c) for c in codes):
                    sys.exit(f"{key} Q{it['no']}: condition '{it['when']}' is not like 'O1, O2'")
                q["showIf"] = {"q": ids[it["parent"]], "in": codes}
            items.append(q)
        forms[key] = {"key": key, "opening": text_of("opening"), "closing": text_of("closing"), "sections": sections, "items": items}

    head = ("// Made by tools/import_excel.py from the team's Excel files - do not edit by hand.\n"
            "// Languages in this file: " + ", ".join(["en"] + sorted(langs)) + ". Missing text falls back to English.\n")
    OUT.write_text(head + "export const FORMS = " + json.dumps(forms, ensure_ascii=False, indent=1) + ";\n", encoding="utf-8")
    for w in warnings:
        print("! " + w)
    for k, f in forms.items():
        print(f"{k}: {len(f['items'])} questions, {len(f['sections'])} sections, conditional: {sum('showIf' in q for q in f['items'])}")
    print("languages:", ", ".join(["en"] + sorted(langs)), "->", OUT.name)


if __name__ == "__main__":
    main()
