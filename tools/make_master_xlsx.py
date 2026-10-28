"""Master translation workbook: one tab per language and form (e.g. "Tamil - Youth"), laid out like the
team's Excel survey tools, with exactly what the live form shows.

    node tools/export_master.mjs > master.json
    python tools/make_master_xlsx.py master.json <output.xlsx>
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

src, out = sys.argv[1], sys.argv[2]
d = json.load(open(src, encoding="utf-8"))
ENGLISH_NAME = {"en": "English", "hi": "Hindi", "kn": "Kannada", "te": "Telugu", "or": "Odia", "ta": "Tamil", "pa": "Punjabi", "mr": "Marathi"}
FORM_NAME = {"youth": "Youth", "parent": "Parent", "public": "General Public"}
SOURCE = {"state": "Dropdown: 36 states / UTs", "district": "Dropdown: districts of the chosen state"}
HEAD = ["Q No.", "Section", "Main Question", "Response Type", "Asked after Q", "Ask if response is"] + [f"Option {i}" for i in range(1, 11)]
FONT = "Nirmala UI"   # Windows font with every Indian script
purple = PatternFill("solid", fgColor="382280")
soft = PatternFill("solid", fgColor="EFEAF8")
wrap = Alignment(wrap_text=True, vertical="top")


def who(code):
    if code == "en":
        return "Team's English survey tools (master)"
    if code == "mr":
        return "Team's Marathi translation (questions); screen text and follow-ups are drafts"
    return "AI draft translation - needs native-speaker review" if code in d["drafts"] else "-"


wb = Workbook()
rm = wb.active
rm.title = "Read me"
rm.append(["1MV Listening Form V3 (#TumBolo) - questions and options in every language"])
rm.append(["Same text as the live form: https://gaupalemunna76.github.io/1mv-listening-form-v3/"])
rm.append(["Option codes O1, O2 ... are used in 'Ask if response is'. Answers are saved in English in the Sheet."])
rm.append([])
rm.append(["Tab", "Language", "Form", "Translation"])
hdr_row = rm.max_row
for l in d["langs"]:
    for key in d["forms"]:
        rm.append([f"{ENGLISH_NAME[l['code']]} - {FORM_NAME[key]}", f"{ENGLISH_NAME[l['code']]} ({l['name']})", FORM_NAME[key], who(l["code"])])
rm["A1"].font = Font(name=FONT, bold=True, size=14, color="382280")
for c in rm[hdr_row]:
    c.font = Font(name=FONT, bold=True, color="FFFFFF"); c.fill = purple
for w, col in zip([26, 22, 16, 70], "ABCD"):
    rm.column_dimensions[col].width = w

for l in d["langs"]:
    code = l["code"]
    t = lambda e: (e or {}).get(code) or (e or {}).get("en") or ""
    for key, f in d["forms"].items():
        ws = wb.create_sheet(f"{ENGLISH_NAME[code]} - {FORM_NAME[key]}")
        ws.append([f"Language: {ENGLISH_NAME[code]} ({l['name']})  |  Form: {FORM_NAME[key]}  |  {who(code)}"])
        ws.append(HEAD)
        ws.append(["", "Opening", t(f["opening"])])
        for r in f["rows"]:
            q = t(r["label"])
            if r["lead"]:
                q = t(r["lead"]) + "\n\n" + q
            opts = [SOURCE[r["source"]]] if r["source"] else [t(o) for o in r["options"]]
            ws.append([r["no"], t(r["section"]), q, r["type"], r["parent"], r["when"]] + opts)
        ws.append(["", "Closing", t(f["closing"])])
        ws["A1"].font = Font(name=FONT, bold=True, size=12, color="382280")
        for c in ws[2]:
            c.font = Font(name=FONT, bold=True, color="FFFFFF"); c.fill = purple; c.alignment = wrap
        for row in ws.iter_rows(min_row=3):
            for c in row:
                c.font = Font(name=FONT, size=11); c.alignment = wrap
        for r in (3, ws.max_row):
            for c in ws[r]:
                c.fill = soft
        for i, w in enumerate([8, 22, 60, 14, 10, 14] + [22] * 10, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "C3"
wb.save(out)
print("tabs:", len(wb.sheetnames), "->", out)
